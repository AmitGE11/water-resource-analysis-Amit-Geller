#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import pandas as pd
import matplotlib.pyplot as plt
from datetime import datetime
import os
import gc
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
import json
from typing import Dict, List, Tuple, Optional
import logging

# Configure logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

class WaterStationAnalyzer:
    """
    Professional water station analyzer that can handle multiple stations
    through configuration-based approach.
    """
    
    def __init__(self, config_file: str = "station_config.json"):
        """
        Initialize the analyzer with station configuration.
        
        Args:
            config_file: Path to JSON configuration file containing station definitions
        """
        self.config = self._load_config(config_file)
        self.stations = self.config.get('stations', {})
        self.base_output_dir = self.config.get('base_output_dir', '')
        self.rain_data_dir = self.config.get('rain_data_dir', '')
        
    def _load_config(self, config_file: str) -> Dict:
        """Load station configuration from JSON file."""
        try:
            with open(config_file, 'r', encoding='utf-8') as f:
                return json.load(f)
        except FileNotFoundError:
            logger.warning(f"Config file {config_file} not found. Using default configuration.")
            return self._get_default_config()
    
    def _get_default_config(self) -> Dict:
        """Return default configuration for stations."""
        return {
            "base_output_dir": r"C:\Users\AmitGeller\Desktop\Yaron Geller\Water_Analysis\Final Analysis",
            "rain_data_dir": r"C:\Users\AmitGeller\Desktop\Yaron Geller\מדידות גשם",
            "stations": {
                "hans_moller": {
                    "name": "Hans Moller",
                    "data_dir": r"C:\Users\AmitGeller\Desktop\Yaron Geller\Water_Analysis\הנס מולר",
                    "threshold": 0.5,
                    "enabled": True
                },
                "haroshet_south": {
                    "name": "Haroshet South",
                    "data_dir": r"C:\Users\AmitGeller\Desktop\Yaron Geller\Water_Analysis\חרושת דרום",
                    "threshold": 0.5,
                    "enabled": True
                },
                "haroshet_north": {
                    "name": "Haroshet North",
                    "data_dir": r"C:\Users\AmitGeller\Desktop\Yaron Geller\Water_Analysis\חרושת צפון",
                    "threshold": 0.5,
                    "enabled": True
                },
                "yosef_levi": {
                    "name": "Yosef Levi",
                    "data_dir": r"C:\Users\AmitGeller\Desktop\Yaron Geller\Water_Analysis\יוסף לוי",
                    "threshold": 0.5,
                    "enabled": True
                },
                "yosef_levi_west": {
                    "name": "Yosef Levi West",
                    "data_dir": r"C:\Users\AmitGeller\Desktop\Yaron Geller\Water_Analysis\יוסף לוי מערב",
                    "threshold": 0.5,
                    "enabled": True
                },
                "measef_shefaram": {
                    "name": "Measef Shefaram",
                    "data_dir": r"C:\Users\AmitGeller\Desktop\Yaron Geller\Water_Analysis\מאסף שפרעם",
                    "threshold": 0.5,
                    "enabled": True
                },
                "shuhot_etzel_habanim": {
                    "name": "Shuhot Etzel - Habanim",
                    "data_dir": r"C:\Users\AmitGeller\Desktop\Yaron Geller\Water_Analysis\שוחת אצל - הבנים",
                    "threshold": 0.5,
                    "enabled": True
                },
                "kiryat_bialik": {
                    "name": "Kiryat Bialik",
                    "data_dir": r"C:\Users\AmitGeller\Desktop\Yaron Geller\Water_Analysis\שוחת כניסה לתש קדמת ביאליק",
                    "threshold": 0.5,
                    "enabled": True
                }
            }
        }
    
    def process_station_csv(self, station_id: str, station_config: Dict) -> List[Dict]:
        """
        Process CSV files for a specific station.
        
        Args:
            station_id: Unique identifier for the station
            station_config: Configuration dictionary for the station
            
        Returns:
            List of processed monthly data dictionaries
        """
        directory_path = station_config['data_dir']
        threshold = station_config['threshold']
        
        logger.info(f"Processing station: {station_config['name']} from {directory_path}")
        
        # Get all CSV files in the directory
        try:
            csv_files = [f for f in os.listdir(directory_path) if f.endswith('.csv')]
            csv_files.sort()
        except FileNotFoundError:
            logger.error(f"Directory not found: {directory_path}")
            return []
        
        logger.info(f"Found {len(csv_files)} CSV files for {station_config['name']}")
        
        monthly_data = []
        
        for csv_file in csv_files:
            file_path = os.path.join(directory_path, csv_file)
            logger.info(f"Processing {csv_file}...")
            
            try:
                df = pd.read_csv(file_path)
                logger.info(f"  - Rows: {len(df)}")
                
                # Convert Time column to string and combine Date and Time columns
                df['Time'] = df['Time'].astype(str)
                df['DateTime'] = pd.to_datetime(df['Date'] + ' ' + df['Time'])
                
                # Find the level column
                level_column = self._find_level_column(df, csv_file)
                if level_column is None:
                    continue
                
                # Clean the data
                df = self._clean_water_data(df, level_column)
                
                # Get month name from filename
                month_name = csv_file.replace('.csv', '')
                
                # Store processed data
                monthly_data.append({
                    'station_id': station_id,
                    'station_name': station_config['name'],
                    'month_name': month_name,
                    'dataframe': df,
                    'level_column': level_column,
                    'threshold': threshold
                })
                
            except Exception as e:
                logger.error(f"Error processing {csv_file}: {str(e)}")
                continue
        
        return monthly_data
    
    def _find_level_column(self, df: pd.DataFrame, filename: str) -> Optional[str]:
        """Find the water level column in the dataframe."""
        # First try to find column with 'level' in name
        for col in df.columns:
            if 'level' in col.lower():
                return col
        
        # If not found, use the last numeric column
        numeric_columns = df.select_dtypes(include=['float64', 'int64']).columns
        if len(numeric_columns) > 0:
            level_column = numeric_columns[-1]
            logger.info(f"Using column '{level_column}' as level measurement for {filename}")
            return level_column
        
        logger.error(f"Could not identify water level column for {filename}")
        return None
    
    def _clean_water_data(self, df: pd.DataFrame, level_column: str) -> pd.DataFrame:
        """Clean water level data by removing invalid values."""
        # Convert to numeric, replacing non-numeric values with NaN
        df[level_column] = pd.to_numeric(df[level_column], errors='coerce')
        
        # Remove NaN values
        df = df.dropna(subset=[level_column])
        
        # Remove negative values and unreasonably large values
        df = df[df[level_column] >= 0]
        df = df[df[level_column] < 10]
        
        return df
    
    def process_rain_data(self) -> Dict:
        """Process rain measurement data for all stations."""
        logger.info(f"Processing rain measurement data from: {self.rain_data_dir}")
        
        rain_data = {}
        
        for station in ['Afek', 'Haifa']:
            csv_file = f"{station}.csv"
            file_path = os.path.join(self.rain_data_dir, csv_file)
            
            try:
                df = pd.read_csv(file_path)
                logger.info(f"Processed {csv_file}: {len(df)} rows")
                
                # Drop unnecessary columns
                columns_to_drop = ['Station']
                if 'code' in df.columns:
                    columns_to_drop.append('code')
                elif 'Code' in df.columns:
                    columns_to_drop.append('Code')
                
                df = df.drop(columns=columns_to_drop)
                rain_data[station] = df
                
            except Exception as e:
                logger.error(f"Error processing {csv_file}: {str(e)}")
                continue
        
        return rain_data
    
    def create_water_level_chart(self, monthly_info: Dict, rain_data: Dict) -> Tuple[str, str]:
        """Create water level charts for a specific month and station."""
        df = monthly_info['dataframe']
        level_column = monthly_info['level_column']
        month_name = monthly_info['month_name']
        station_name = monthly_info['station_name']
        threshold = monthly_info['threshold']
        
        # Create high level DataFrame
        high_level_df = df[df[level_column] > threshold]
        
        # Create output directory for this station
        station_output_dir = os.path.join(self.base_output_dir, station_name.replace(' ', '_'))
        os.makedirs(station_output_dir, exist_ok=True)
        
        # Create both dots and bars charts
        chart_paths = self._create_chart_versions(
            df, high_level_df, level_column, month_name, station_name, 
            threshold, station_output_dir, rain_data
        )
        
        return chart_paths
    
    def _create_chart_versions(self, df: pd.DataFrame, high_level_df: pd.DataFrame, 
                              level_column: str, month_name: str, station_name: str,
                              threshold: float, output_dir: str, rain_data: Dict) -> Tuple[str, str]:
        """Create both dots and bars versions of the water level chart."""
        
        # Version 1: Dots
        fig1, ax1 = plt.subplots(figsize=(15, 8))
        self._plot_water_levels(ax1, df, high_level_df, level_column, threshold, rain_data, 'scatter')
        plt.title(f'{station_name} - Water Level and Rainfall Measurements - {month_name} (Dots)')
        plt.tight_layout()
        
        dots_path = os.path.join(output_dir, f'water_level_chart_{month_name}_dots.png')
        plt.savefig(dots_path, dpi=300, bbox_inches='tight')
        plt.close(fig1)
        
        # Version 2: Bars
        fig2, ax2 = plt.subplots(figsize=(15, 8))
        self._plot_water_levels(ax2, df, high_level_df, level_column, threshold, rain_data, 'bar')
        plt.title(f'{station_name} - Water Level and Rainfall Measurements - {month_name} (Bars)')
        plt.tight_layout()
        
        bars_path = os.path.join(output_dir, f'water_level_chart_{month_name}_bars.png')
        plt.savefig(bars_path, dpi=300, bbox_inches='tight')
        plt.close(fig2)
        
        return dots_path, bars_path
    
    def _plot_water_levels(self, ax, df: pd.DataFrame, high_level_df: pd.DataFrame,
                          level_column: str, threshold: float, rain_data: Dict, plot_type: str):
        """Plot water levels and rain data on the given axis."""
        
        # Plot water levels on primary y-axis
        ax.plot(df['DateTime'], df[level_column], color='blue', alpha=0.5, label='Water Level')
        
        # Highlight high water levels
        ax.scatter(high_level_df['DateTime'], high_level_df[level_column],
                  color='red', label=f'High Water Level (>{threshold}m)')
        
        # Add threshold line
        ax.axhline(y=threshold, color='r', linestyle='--', alpha=0.3, label=f'{threshold}m Threshold')
        
        # Customize primary y-axis
        ax.set_xlabel('Date')
        ax.set_ylabel('Water Level (m)', color='blue')
        ax.tick_params(axis='y', labelcolor='blue')
        ax.set_ylim(0, max(df[level_column].max() * 1.1, threshold + 0.1))
        ax.grid(True, alpha=0.3)
        
        # Create secondary y-axis for rain data
        ax2 = ax.twinx()
        
        # Plot rain data
        self._plot_rain_data(ax2, df, rain_data, plot_type)
        
        # Combine legends
        lines1, labels1 = ax.get_legend_handles_labels()
        lines2, labels2 = ax2.get_legend_handles_labels()
        ax.legend(lines1 + lines2, labels1 + labels2, loc='upper left')
        
        # Rotate x-axis labels
        plt.xticks(rotation=45)
    
    def _plot_rain_data(self, ax, df: pd.DataFrame, rain_data: Dict, plot_type: str):
        """Plot rain data on the secondary axis."""
        for station, rain_df in rain_data.items():
            if len(rain_df) > 0:
                rain_df_copy = rain_df.copy()
                rain_df_copy['Date'] = pd.to_datetime(rain_df_copy['Date'], format='%d/%m/%Y')
                
                # Filter rain data to match the month period
                start_date = df['DateTime'].min().date()
                end_date = df['DateTime'].max().date()
                
                month_rain = rain_df_copy[
                    (rain_df_copy['Date'].dt.date >= start_date) & 
                    (rain_df_copy['Date'].dt.date <= end_date)
                ]
                
                if len(month_rain) > 0:
                    color = 'green' if station == 'Afek' else 'orange'
                    
                    if plot_type == 'scatter':
                        ax.scatter(month_rain['Date'], month_rain['Rain'], 
                                 color=color, s=100, edgecolor='black', linewidth=1,
                                 label=f'{station} Rain (mm)')
                    else:  # bar
                        bar_width = 0.35
                        if station == 'Afek':
                            ax.bar([d - pd.Timedelta(hours=6) for d in month_rain['Date']], 
                                  month_rain['Rain'], color=color, alpha=0.7, width=bar_width,
                                  label=f'{station} Rain (mm)')
                        else:
                            ax.bar([d + pd.Timedelta(hours=6) for d in month_rain['Date']], 
                                  month_rain['Rain'], color=color, alpha=0.7, width=bar_width,
                                  label=f'{station} Rain (mm)')
        
        ax.set_ylabel('Rainfall (mm)', color='green')
        ax.tick_params(axis='y', labelcolor='green')
    
    def export_station_to_excel(self, station_monthly_data: List[Dict], chart_paths: List[Tuple[str, str]]) -> str:
        """Export all data for a station to Excel."""
        if not station_monthly_data:
            return ""
        
        station_name = station_monthly_data[0]['station_name']
        station_output_dir = os.path.join(self.base_output_dir, station_name.replace(' ', '_'))
        excel_file = os.path.join(station_output_dir, f'{station_name}_Analysis.xlsx')
        
        logger.info(f"Creating Excel file: {excel_file}")
        
        with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
            for i, month_info in enumerate(station_monthly_data):
                df = month_info['dataframe']
                level_column = month_info['level_column']
                month_name = month_info['month_name']
                threshold = month_info['threshold']
                
                # Create high level DataFrame
                high_level_df = df[df[level_column] > threshold]
                
                # Calculate statistics
                num_high_readings = len(high_level_df)
                total_minutes = num_high_readings * 5
                hours = total_minutes // 60
                remaining_minutes = total_minutes % 60
                
                stats_dict = {
                    'Month': month_name,
                    'Total Readings': len(df),
                    f'High Water Level Readings (>{threshold}m)': num_high_readings,
                    f'Total Hours Above {threshold}m': hours,
                    f'Additional Minutes Above {threshold}m': remaining_minutes,
                    'Maximum Water Level': df[level_column].max(),
                    'Average Water Level': df[level_column].mean(),
                }
                
                # Export data to sheets
                df.to_excel(writer, sheet_name=f"{month_name}_All_Measurements", index=False)
                high_level_df.to_excel(writer, sheet_name=f"{month_name}_High_Levels", index=False)
                stats_df = pd.DataFrame([stats_dict])
                stats_df.to_excel(writer, sheet_name=f"{month_name}_Statistics", index=False)
                
                # Add charts if available
                if i < len(chart_paths):
                    self._add_charts_to_excel(writer, chart_paths[i], month_name)
        
        return excel_file
    
    def _add_charts_to_excel(self, writer, chart_paths: Tuple[str, str], month_name: str):
        """Add charts to Excel sheet."""
        workbook = writer.book
        chart_sheet = workbook.create_sheet(f"{month_name}_Charts")
        
        # Add dots chart
        if os.path.exists(chart_paths[0]):
            img_dots = Image(chart_paths[0])
            img_dots.width = 800
            img_dots.height = 400
            chart_sheet.add_image(img_dots, 'A1')
        
        # Add bars chart
        if os.path.exists(chart_paths[1]):
            img_bars = Image(chart_paths[1])
            img_bars.width = 800
            img_bars.height = 400
            chart_sheet.add_image(img_bars, 'A25')
    
    def analyze_all_stations(self):
        """Analyze all enabled stations."""
        logger.info("Starting analysis of all stations...")
        
        # Process rain data once for all stations
        rain_data = self.process_rain_data()
        
        all_results = {}
        
        for station_id, station_config in self.stations.items():
            if not station_config.get('enabled', True):
                logger.info(f"Skipping disabled station: {station_config['name']}")
                continue
            
            try:
                logger.info(f"\n{'='*50}")
                logger.info(f"Processing station: {station_config['name']}")
                logger.info(f"{'='*50}")
                
                # Process station data
                monthly_data = self.process_station_csv(station_id, station_config)
                
                if not monthly_data:
                    logger.warning(f"No data found for station: {station_config['name']}")
                    continue
                
                # Create charts for each month
                chart_paths = []
                for month_info in monthly_data:
                    dots_path, bars_path = self.create_water_level_chart(month_info, rain_data)
                    chart_paths.append((dots_path, bars_path))
                    
                    # Print summary statistics
                    df = month_info['dataframe']
                    level_column = month_info['level_column']
                    threshold = month_info['threshold']
                    high_level_df = df[df[level_column] > threshold]
                    
                    num_high_readings = len(high_level_df)
                    total_minutes = num_high_readings * 5
                    hours = total_minutes // 60
                    remaining_minutes = total_minutes % 60
                    
                    logger.info(f"  {month_info['month_name']}: {len(df)} readings, "
                              f"{num_high_readings} high readings, "
                              f"{hours}h {remaining_minutes}m above {threshold}m")
                
                # Export to Excel
                excel_file = self.export_station_to_excel(monthly_data, chart_paths)
                
                # Store results
                all_results[station_id] = {
                    'station_name': station_config['name'],
                    'monthly_data': monthly_data,
                    'chart_paths': chart_paths,
                    'excel_file': excel_file
                }
                
                # Clean up temporary chart files
                for chart_path_tuple in chart_paths:
                    try:
                        os.remove(chart_path_tuple[0])
                        os.remove(chart_path_tuple[1])
                    except:
                        pass
                
                logger.info(f"Completed analysis for {station_config['name']}")
                
            except Exception as e:
                logger.error(f"Error processing station {station_config['name']}: {str(e)}")
                continue
        
        logger.info(f"\n{'='*50}")
        logger.info("Analysis completed!")
        logger.info(f"Successfully processed {len(all_results)} stations")
        logger.info(f"{'='*50}")
        
        return all_results

def main():
    """Main function to run the water station analysis."""
    analyzer = WaterStationAnalyzer()
    results = analyzer.analyze_all_stations()
    
    # Print summary
    print("\n" + "="*60)
    print("ANALYSIS SUMMARY")
    print("="*60)
    for station_id, result in results.items():
        print(f"\n{result['station_name']}:")
        print(f"  Excel file: {result['excel_file']}")
        print(f"  Months processed: {len(result['monthly_data'])}")

if __name__ == "__main__":
    main()
